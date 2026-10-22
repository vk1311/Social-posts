from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

# (audience, town, place, address, type, leave, ask, hours, priority)
W,A,B="Waypoint","Apps","Both"
WF="Waypoint flyer on board"; WC="Waypoint counter cards (small stack)"
AP="Apps poster on board"; AC="Apps take-one cards (small stack)"
BOTH_LEAVE="Waypoint counter cards + Apps take-one cards"
rows=[
# Kentville downtown
(B,"Kentville","CBDC Hants-Kings","448 Main St, Kentville, NS B4N 1R5","Small-business lender",WC+" + "+AC,"They lend to and advise new owners. Ask to leave cards for clients and if you can drop off more later.","Mon–Fri 8:30–4:30",1),
(B,"Kentville","Valley Regional Enterprise Network (Valley REN)","448 Main St, Kentville, NS B4N 1R5","Business support office",BOTH_LEAVE,"Same building as CBDC. Ask about their business newsletter and leaving cards.","Mon–Fri 8:30–4:30",1),
(B,"Kentville","Annapolis Valley Chamber of Commerce","448 Main St, Kentville, NS B4N 1R5","Business support office",BOTH_LEAVE,"Ask about the members' board, the newsletter, and leaving cards.","Mon–Fri 8:30–4",1),
(B,"Kentville","The Space and Company (coworking)","10 Webster St #210, Kentville, NS","Coworking",WF+" + "+AC,"Freelancers and small owners work here. Ask to leave cards in the common area.","Weekdays",1),
(W,"Kentville","mbf Chartered Professional Accountants","464 Main St, Kentville, NS B4N 1K8","Accountant",WC,"Ask if you can leave a few cards for clients who are still on paper. Don't pitch the apps here.","Mon–Fri 8:30–4:30",2),
(W,"Kentville","Infinite Solutions Bookkeeping & Accounting","25 Webster Ct, Kentville, NS B4N 1H3","Bookkeeper",WC,"Same ask as accountants: cards for clients drowning in paperwork.","Mon–Fri 8:30–4:30",2),
(W,"Kentville","Rockwell Home Hardware","414 Main St, Kentville, NS","Hardware / trades",WF+" + "+WC,"Trades buy here. Ask about a board by the door or cards at the till.","Daily",1),
(A,"Kentville","Kentville Library","440 Main St, Kentville, NS","Library board",AP,"Community board inside. Ask staff to approve and date-stamp.","Tue–Sat (closed Sun, Mon)",1),
(A,"Kentville","Kentville Farmers Market","Centre Square, Kentville, NS B4N 1G9","Farmers' market",AC,"Wednesdays. Vendors are the perfect Apps user. Ask the market manager first, then chat to vendors.","Wed 10–1",1),
(A,"Kentville","Wheelhouse Coffee Co.","2 Webster St, Kentville, NS","Café",AC,"Ask at the counter.","Mon–Sat",2),
(A,"Kentville","Half Acre Cafe","395 Main St, Kentville, NS","Café",AC,"Ask at the counter.","Daily 8–4",2),
(A,"Kentville","T.A.N. Coffee Kentville","431 Main St, Kentville, NS","Café",AC,"Ask at the counter.","Daily",2),
(A,"Kentville","Kentville Recreation Centre","348 Main St, Kentville, NS","Community centre",AP,"Ask at the front desk.","Mon–Fri 8:30–4:30",3),
(A,"Kentville","Royal Canadian Legion Branch 6","37 River St, Kentville, NS","Legion hall",AP,"Hall board. Ask the bar or front volunteer.","Thu–Sun",3),
(A,"Kentville","Maytag Coin Laundry","252 Main St, Kentville, NS","Laundromat",AP,"Ask at the Needs store next door.","Daily 9–9",3),
# Kentville west / Coldbrook
(W,"Kentville","Fraser's Pro Home Centre","14 Pelton Dr, Kentville, NS","Building supply",WC,"Contractor desk. Ask to leave cards at the counter.","Mon–Sat",1),
(A,"Kentville","Credit Union Rec Complex","38 Gary Pearl Dr, Kentville, NS","Community centre",AP,"Ask at the front desk.","Weekdays",3),
(W,"Coldbrook","Scotian Gold Country Store","7027 Hwy 1, Coldbrook, NS","Farm / landscape supply",WF+" + "+WC,"Farmers and landscapers shop here. Ask about the community board.","Mon–Sat",1),
(W,"Coldbrook","Carquest Auto Parts","108 Coldbrook Village Park Dr, Coldbrook, NS","Auto parts",WC,"Mechanics and fleet owners. Cards at the counter.","Mon–Sat",2),
(W,"Kentville","Kent & Duffett Chartered Professional Accountants","46 Chipman Dr Unit 1, Kentville, NS B4N 3V7","Accountant",WC,"Cards for clients still doing invoices by hand.","Weekdays",3),
(A,"Kentville","NSCC Kingstec Campus","236 Belcher St, Kentville, NS","College boards",AP,"Student boards. Ask student services for an approval stamp.","Mon–Sat",2),
# New Minas
(W,"New Minas","Valley Power Products (rentals)","9 Nova Dr, New Minas, NS B4N 3G7","Equipment rental",WF+" + "+WC,"Contractors rent here. Ask about a board or cards at the counter.","Mon–Fri 7:30–5",1),
(W,"New Minas","Fleetworx Inc.","64 Minas Warehouse Rd, New Minas, NS B4N 5G3","Truck / fleet",WC,"Fleet owners. Cards at the counter.","Mon–Fri 8–5",2),
(W,"New Minas","Irving Big Stop (New Minas)","5477 Prospect Rd, New Minas, NS B4N 3K9","Truck stop / diner",WF,"Truckers and contractors eat breakfast here. Ask about a board near the restaurant.","24 hours (restaurant 8–8)",2),
(W,"New Minas","Kent Building Supplies","5447 Prospect Rd, New Minas, NS","Building supply",WF+" + "+WC,"Contractor desk. Ask about the board and cards.","Daily",1),
(W,"New Minas","Bumper to Bumper Auto Parts","5494 Prospect Rd, New Minas, NS","Auto parts",WC,"Cards at the counter.","Mon–Sat",2),
(W,"New Minas","The Home Depot","21 Silver Fox Ave, New Minas, NS","Building supply",WC,"Pro desk. Ask whether they allow cards.","Daily",2),
(W,"New Minas","NAPA Auto Parts","8793 Commercial St, New Minas, NS","Auto parts",WC,"Trade counter. Cards at the till.","Mon–Sat",2),
(W,"New Minas","Repracorp Design and Print","8927 Commercial St, New Minas, NS","Print shop",WC,"Business print customers. Ask at the front.","Mon–Fri 8–5",2),
(B,"New Minas","Running Man Grocery Mart","9036 Commercial St, New Minas, NS","Bulk grocery",BOTH_LEAVE,"Restaurant and shop owners stock up here. Ask at the counter.","Daily",2),
(B,"New Minas","The UPS Store","9049 Commercial St, New Minas, NS","Shipping",AC+" + "+WC,"Online sellers and small shops ship here. Ask to leave cards.","Mon–Sat",1),
(A,"New Minas","Atlantic Superstore","9064 Commercial St, New Minas, NS","Grocery board",AP,"Community board near the entrance. Ask customer service.","Daily 7–10",1),
(B,"New Minas","Staples + Staples Print","9081 Commercial St, New Minas, NS","Print shop",BOTH_LEAVE,"Ask at the print desk about leaving cards.","Daily",1),
(A,"New Minas","Walmart Supercentre","9097 Commercial St, New Minas, NS","Store board",AP,"Ask customer service whether there's a community board.","Daily 7–10",2),
(B,"New Minas","Panda Print","9146 Commercial St, New Minas, NS","Print shop",BOTH_LEAVE,"Owners print signs and cards here. Ask to leave cards on the counter.","Mon–Fri 9–5",1),
(W,"New Minas","Valley Credit Union","9237 Commercial St, New Minas, NS","Credit union",WC,"Business banking. Ask whether they have a board or take cards.","Mon–Fri 10–4",2),
(A,"New Minas","Sobeys New Minas","9256 Commercial St, New Minas, NS","Grocery board",AP,"Community board by the doors. Ask customer service.","Daily 7–10",1),
(A,"New Minas","M&M Food Market","9256 Commercial St, New Minas, NS","Store",AC,"Ask at the counter.","Daily",3),
# Greenwich / Wolfville
(W,"Greenwich","TNT Tool Rental & Safety","9776 Hwy 1, Greenwich, NS B4P 0A7","Equipment rental",WF+" + "+WC,"Contractors rent here. Ask about a board or cards at the counter.","Mon–Fri 7:30–5",1),
(A,"Greenwich","Noggins Corner Farm Market","10009 Hwy 1, Greenwich, NS","Farm market",AP,"Local producers sell here. Ask about the board.","Daily 8–8",2),
(A,"Greenwich","Hennigar's Farm Market","10272 Hwy 1, Wolfville, NS","Farm market",AP,"Local producers sell here. Ask about the board.","Daily 8–8",2),
(A,"Wolfville","Wolfville Farmers' Market","24 Elm Ave, Wolfville, NS","Farmers' market",AP+" + "+AC,"Saturday mornings. Ask the market office about the board, then talk to vendors.","Sat 8:30–1",1),
(A,"Wolfville","Wolfville Memorial Library","21 Elm Ave, Wolfville, NS","Library board",AP,"Community board inside. Ask staff.","Tue–Sat",1),
(A,"Wolfville","Charts Café","16 Elm Ave, Wolfville, NS","Café",AC,"Ask at the counter.","Daily",2),
(A,"Wolfville","Just Us! Wolfville Café","450 Main St, Wolfville, NS","Café",AC,"Ask at the counter.","Daily",2),
(A,"Wolfville","T.A.N. Coffee Wolfville","378 Main St, Wolfville, NS","Café",AC,"Ask at the counter.","Daily",2),
(A,"Wolfville","Royal Canadian Legion Branch 74","310 Main St, Wolfville, NS","Legion hall",AP,"Hall board. Ask a volunteer.","Tue–Sat",3),
(A,"Wolfville","Acadia Students' Union","30 Highland Ave, Wolfville, NS","University boards",AP,"Poster rules apply. Ask at the front desk for a stamp.","Daily",2),
# Canning
(A,"Canning","The Village Coffeehouse","9844 Main St, Canning, NS","Café",AC,"Local gathering spot. Ask at the counter.","Mon–Sat",3),
(W,"Canning","Maple Leaf Home Hardware","9767 Main St, Canning, NS","Hardware / trades",WF+" + "+WC,"Ask about a board or the till.","Mon–Sat",2),
# Berwick and west
(A,"Berwick","Berwick and District Library","236 Commercial St, Berwick, NS","Library board",AP,"Community board. Ask staff.","Tue–Sat",2),
(A,"Berwick","Foodland Berwick","223 Commercial St, Berwick, NS","Grocery board",AP,"Board near the entrance.","Daily",2),
(W,"Berwick","Feeds'n Needs","229 Commercial St, Berwick, NS","Farm / pet supply",WF+" + "+WC,"Farm and pet businesses. Ask at the counter.","Daily",2),
(A,"Woodville","Woodville Community Centre","342 Bligh Rd, Woodville, NS","Community centre",AP,"Hall board.","Varies",3),
(W,"Cambridge","Cambridge Construction Rentals","5660 Evangeline Trl, Cambridge, NS B0P 1G0","Equipment rental",WF+" + "+WC,"Contractors rent here. Ask about a board or cards at the counter.","Mon–Fri 7–5",2),
(A,"Cambridge","Cambridge Community Centre","5961 Hwy 1, Cambridge, NS","Community centre",AP,"Hall board.","Varies",3),
(A,"Aylesford","Dempsey Corner Orchards","2717 NS-221, Aylesford, NS","Farm market",AP,"Ask about the board.","Mon–Sat",3),
# Windsor
(A,"Windsor","Windsor Regional Library","195 Albert St, Windsor, NS","Library board",AP,"Community board. Ask staff.","Tue–Sat",2),
(A,"Windsor","Gerrish & Gray Cafe","134 Gerrish St, Windsor, NS","Café",AC,"Ask at the counter.","Daily",3),
(A,"Falmouth","Expedition Coffee Roasters","1813 Hwy 1, Falmouth, NS","Café",AC,"Ask at the counter.","Tue–Wed mornings",3),
# Seasonal
(A,"Valley-wide","Holiday craft fairs (Nov–Dec)","Check novascotia.com holiday markets list + town Facebook pages","Craft fair",AC,"Best fit for Apps right now. Ask organisers if you can leave cards on the info table, or talk to vendors.","Weekends Nov–Dec",1),
]

ink="12252F"; mag="C2185B"; paper="F2ECE0"
hfont=Font(name="Arial",bold=True,color="FFFFFF",size=10)
hfill=PatternFill("solid",fgColor=ink)
bfont=Font(name="Arial",size=10)
thin=Side(style="thin",color="D9D0C0")
border=Border(bottom=thin)
headers=["Stop","Audience","Town","Place","Address (paste into Maps)","Type","What to leave","What to ask","Hours (check first)","Priority (1 = do first)","Done (Y)","Notes"]
widths=[6,10,13,38,44,22,40,60,22,12,9,30]

wb=Workbook()
def sheet(ws,title,data,intro):
    ws.title=title
    ws["A1"]=title; ws["A1"].font=Font(name="Arial",bold=True,size=14,color=ink)
    ws["A2"]=intro; ws["A2"].font=Font(name="Arial",size=10,color="5A6B73")
    ws.merge_cells("A2:L2")
    for i,h in enumerate(headers,1):
        c=ws.cell(row=4,column=i,value=h); c.font=hfont; c.fill=hfill
        c.alignment=Alignment(vertical="center",wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width=widths[i-1]
    ws.row_dimensions[4].height=30
    for r,row in enumerate(data,5):
        aud,town,place,addr,typ,leave,ask,hours,pri=row
        vals=[r-4,aud,town,place,addr,typ,leave,ask,hours,pri,"",""]
        for i,v in enumerate(vals,1):
            c=ws.cell(row=r,column=i,value=v); c.font=bfont; c.border=border
            c.alignment=Alignment(vertical="top",wrap_text=i in (4,5,7,8,12))
        if pri==1:
            ws.cell(row=r,column=10).font=Font(name="Arial",size=10,bold=True,color=mag)
        ws.cell(row=r,column=11).fill=PatternFill("solid",fgColor="FFF7CC")
    last=4+len(data)
    dv=DataValidation(type="list",formula1='"Y,"',allow_blank=True); ws.add_data_validation(dv); dv.add(f"K5:K{last}")
    ws.freeze_panes="E5"
    ws.auto_filter.ref=f"A4:L{last}"
    # done counter
    ws["J3"]="Done:"; ws["J3"].font=Font(name="Arial",bold=True,size=10); ws["J3"].alignment=Alignment(horizontal="right")
    ws["K3"]=f'=COUNTIF(K5:K{last},"Y")&" of "&COUNTA(D5:D{last})'
    ws["K3"].font=Font(name="Arial",bold=True,size=10,color=mag)
    return last

wp=[r for r in rows if r[0] in (W,B)]
ap=[r for r in rows if r[0] in (A,B)]
ws1=wb.active
sheet(ws1,"Waypoint flyers",wp,"Business owners with paperwork pain: trades counters, rentals, lenders, accountants, print shops. Leave the Waypoint flyer on boards and counter cards at tills. Yellow column: type Y when done.")
ws2=wb.create_sheet()
sheet(ws2,"Waypoint Apps",ap,"Solo sellers, makers, side hustles, students: markets, libraries, grocery boards, cafés. Leave the Apps poster on boards and take-one cards at counters. Yellow column: type Y when done.")
ws3=wb.create_sheet()
allrows=sorted(rows,key=lambda r:(["Kentville","Coldbrook","New Minas","Greenwich","Wolfville","Canning","Berwick","Woodville","Cambridge","Aylesford","Windsor","Falmouth","Valley-wide"].index(r[1])))
last3=sheet(ws3,"All stops (route)",allrows,"Every stop once, grouped by town in driving order (Kentville → New Minas → Wolfville → Canning → west → Windsor). Copy the Address column into your route planner with your Keegix stops.")

# Print list
ws4=wb.create_sheet("Print list")
ws4["A1"]="Print list"; ws4["A1"].font=Font(name="Arial",bold=True,size=14,color=ink)
ws4["A2"]="Counts update from the 'All stops (route)' sheet. Change the blue numbers for more or fewer per stop, or more spares. Page numbers are in waypoint-flyers-print.pdf."; ws4["A2"].font=Font(name="Arial",size=10,color="5A6B73")
hdr=["Item","Stops that need it","Per stop","Spares","Pieces","Per sheet","Sheets to print","Paper","Sides","Notes"]
for i,h in enumerate(hdr,1):
    c=ws4.cell(row=4,column=i,value=h); c.font=hfont; c.fill=hfill; c.alignment=Alignment(vertical="center",wrap_text=True)
for i,w in enumerate([36,12,10,9,10,10,12,12,28,52],1): ws4.column_dimensions[get_column_letter(i)].width=w
ws4.row_dimensions[4].height=30
# (item, stop-list match, per stop, spares, per sheet, paper, sides, notes)
items=[("Waypoint flyer A or B (Letter)","*Waypoint flyer*",1,2,1,"Paper","One-sided","Plain copy paper. PDF page 1 (A) or 2 (B)."),
       ("Waypoint counter cards (quarter-letter)","*Waypoint counter cards*",6,0,4,"Cardstock","Two-sided, flip on long edge","PDF pages 3 (front) and 4 (back). Cut on the dashed lines into 4."),
       ("Apps poster A or B (Letter)","*Apps poster*",1,0,1,"Paper","One-sided","Paper is fine on boards. PDF page 5 (A) or 6 (B)."),
       ("Apps take-one cards (quarter-letter)","*Apps take-one cards*",8,0,4,"Cardstock","One-sided","Cut on the crop marks. PDF page 7.")]
blue=Font(name="Arial",size=10,color="0000FF")
for r,(name,pat,per,spare,ps,paper,sides,note) in enumerate(items,5):
    ws4.cell(row=r,column=1,value=name).font=bfont
    ws4.cell(row=r,column=2,value=f"=COUNTIF('All stops (route)'!G5:G{last3},\"{pat}\")").font=bfont
    c=ws4.cell(row=r,column=3,value=per); c.font=blue
    c=ws4.cell(row=r,column=4,value=spare); c.font=blue
    ws4.cell(row=r,column=5,value=f"=B{r}*C{r}+D{r}").font=bfont
    c=ws4.cell(row=r,column=6,value=ps); c.font=blue
    ws4.cell(row=r,column=7,value=f"=ROUNDUP(E{r}/F{r},0)").font=Font(name="Arial",size=10,bold=True,color=mag)
    ws4.cell(row=r,column=8,value=paper).font=bfont
    ws4.cell(row=r,column=9,value=sides).font=bfont
    c=ws4.cell(row=r,column=10,value=note); c.font=bfont; c.alignment=Alignment(wrap_text=True,vertical="top")
for r,(label,f) in enumerate([("Cardstock sheets (65 lb)",'=SUMIF(H5:H8,"Cardstock",G5:G8)'),("Copy paper sheets",'=SUMIF(H5:H8,"Paper",G5:G8)'),("Total sheets","=G10+G11")],10):
    ws4.cell(row=r,column=1,value=label).font=Font(name="Arial",bold=True,size=10)
    ws4.cell(row=r,column=7,value=f).font=Font(name="Arial",bold=True,size=10,color=mag)
ws4["A14"]="Bring"; ws4["A14"].font=Font(name="Arial",bold=True,size=10)
for i,t in enumerate(["Pushpins and a roll of tape","Scissors or a paper trimmer for the cards","Your phone with app.waypointns.ca open, to show someone in 30 seconds","Ask first everywhere. Boards get cleared every 2–4 weeks, so re-post on the 1st of each month."],15):
    ws4.cell(row=i,column=1,value="• "+t).font=bfont
ws4["A20"]="Counter cards print on both sides: see 'How to print two-sided' in poster-run/README.md. Print one test sheet first."; ws4["A20"].font=Font(name="Arial",size=10,bold=True,color=ink)
ws4["A21"]="Addresses and hours come from Google Maps and earlier route notes (Oct 2026). Check hours before a long drive."; ws4["A21"].font=Font(name="Arial",size=9,italic=True,color="5A6B73")

import os
out=os.path.join(os.path.dirname(os.path.abspath(__file__)),"..","waypoint-flyer-run.xlsx")
wb.save(out); print(out, len(wp), len(ap), len(rows))
